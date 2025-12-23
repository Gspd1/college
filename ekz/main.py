# main.py
import os
from docx import Document
from docx.shared import Inches, Pt
from tkinter import filedialog
import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
from datetime import date, datetime, timedelta
import sys
from db import *
from create_db import create_database

# Сначала создаем БД если нужно
create_database()

page_names = ["Главная", "Запись", "График работы", "Услуги", "Клиенты", "Мастера", "Отчеты"]
pages = {}
menu_buttons = {}
current_page = None

# Глобальные переменные для хранения состояния
selected_time_slot = None  # Выбранный временной слот на странице записи
selected_master_id = None  # Выбранный мастер на странице записи

root = tk.Tk()
root.title("Система управления салоном красоты")
root.state('zoomed')
root.configure(bg="#F0F0F0")


def quit_app():
    """Выход из приложения"""
    root.quit()
    root.destroy()


def show_page(page_name):
    """Переключение между страницами"""
    global current_page
    if current_page and current_page in menu_buttons:
        menu_buttons[current_page].configure(style="Menu.TButton")
    pages[page_name].tkraise()
    current_page = page_name
    menu_buttons[page_name].configure(style="Selected.Menu.TButton")

    # Обновляем данные при переходе на страницу
    if page_name == "Клиенты":
        refresh_clients_page()
    elif page_name == "Мастера":
        refresh_masters_page()
    elif page_name == "Услуги":
        refresh_services_page()
    elif page_name == "Запись":
        refresh_booking_page()
    elif page_name == "График работы":
        refresh_schedule_page()
    elif page_name == "Главная":
        refresh_main_page()


# --- Стили ---
LARGE_FONT = ('Arial', 20)
MEDIUM_FONT = ('Arial', 16)
SMALL_FONT = ('Arial', 14)
BIG_TIME_FONT = ('Arial', 18, 'bold')

style = ttk.Style()
style.configure("Menu.TButton", font=LARGE_FONT, padding=12)
style.configure("Selected.Menu.TButton", background="#CCCCCC", font=LARGE_FONT, padding=12)
style.configure("Large.TButton", font=MEDIUM_FONT, padding=8)
style.configure("Treeview", font=MEDIUM_FONT)
style.configure("Treeview.Heading", font=MEDIUM_FONT)
style.configure("TLabel", font=MEDIUM_FONT)
style.configure("TEntry", font=MEDIUM_FONT)


def make_large_entry(widget):
    """Настройка шрифта для полей ввода"""
    try:
        widget.configure(font=MEDIUM_FONT)
    except Exception:
        pass


def make_large_combo(combo):
    """Настройка шрифта для комбобоксов"""
    try:
        combo.configure(font=MEDIUM_FONT, width=16)
    except Exception:
        try:
            combo.configure(width=16)
        except Exception:
            pass


# ========== ГЛАВНАЯ СТРАНИЦА ==========
def refresh_main_page():
    """Обновление главной страницы"""
    page = pages["Главная"]

    # Очищаем старые данные в таблице
    for item in main_tree.get_children():
        main_tree.delete(item)

    # Получаем сегодняшние записи
    today_str = date.today().strftime('%Y-%m-%d')
    try:
        appointments = get_appointments_by_date(today_str)
    except Exception as e:
        print(f"Ошибка при загрузке записей: {e}")
        appointments = []

    # Добавляем записи в таблицу
    for app in appointments:
        # Проверяем наличие всех необходимых полей
        start_time = app.get('start_time', '')
        if start_time:
            if isinstance(start_time, str):
                start_time_str = start_time[:5]  # Берем только часы и минуты
            else:
                start_time_str = start_time.strftime('%H:%M')
        else:
            start_time_str = ''

        master_name = app.get('master_name', 'Не указан')
        service_name = app.get('service_name', 'Не указана')
        client_name = app.get('client_name', 'Не указан')

        main_tree.insert("", tk.END, values=(
            start_time_str,
            master_name,
            service_name,
            client_name
        ))


def create_main_page(page):
    """Создание главной страницы"""
    page.grid_rowconfigure(4, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="Главная", font=LARGE_FONT).grid(row=0, column=0, pady=10, sticky="n")
    ttk.Label(page, text=f"Сегодня: {date.today().strftime('%d.%m.%Y')}", font=MEDIUM_FONT).grid(
        row=1, column=0, sticky="nw", padx=12, pady=5
    )

    # Статистика
    stats_frame = ttk.Frame(page)
    stats_frame.grid(row=2, column=0, sticky="ew", padx=12, pady=(6, 12))
    stats_frame.grid_columnconfigure(0, weight=1)

    ttk.Label(stats_frame, text="Текущие показатели", font=MEDIUM_FONT).grid(row=0, column=0, pady=6, sticky="w")

    # Получаем статистику за сегодня
    today_str = date.today().strftime('%Y-%m-%d')
    try:
        appointments = get_appointments_by_date(today_str)
    except Exception as e:
        print(f"Ошибка при загрузке статистики: {e}")
        appointments = []

    total = len(appointments)
    completed = len([a for a in appointments if a.get('status') == 'Выполнено'])
    cancelled = len([a for a in appointments if a.get('status') == 'Отменено'])
    pending = len([a for a in appointments if a.get('status') == 'Ожидается'])
    in_progress = len([a for a in appointments if a.get('status') == 'В процессе'])

    row1_bg = "#D3D3D3"
    row2_bg = "#D3D3D3"

    first_row = tk.Frame(stats_frame, bg=row1_bg)
    first_row.grid(row=1, column=0, sticky="ew", padx=20, pady=(2, 0))
    for c in range(3):
        first_row.grid_columnconfigure(c, weight=1, uniform="statscol")

    second_row = tk.Frame(stats_frame, bg=row2_bg)
    second_row.grid(row=2, column=0, sticky="ew", padx=20, pady=(8, 6))
    for c in range(3):
        second_row.grid_columnconfigure(c, weight=1, uniform="statscol")

    stats_data1 = [
        f"Записей сегодня: {total}",
        f"Выполнено: {completed}",
        f"Отменено: {cancelled}"
    ]

    stats_data2 = [
        f"Ожидается: {pending}",
        f"В работе: {in_progress}",
        ""
    ]

    for col, text in enumerate(stats_data1):
        lbl = tk.Label(first_row, text=text, font=SMALL_FONT, bg=row1_bg, anchor="center")
        lbl.grid(row=0, column=col, sticky="nsew", padx=6, pady=8)

    for col, text in enumerate(stats_data2):
        lbl = tk.Label(second_row, text=text, font=SMALL_FONT, bg=row2_bg, anchor="center")
        lbl.grid(row=0, column=col, sticky="nsew", padx=6, pady=8)

    ttk.Label(page, text="Ближайшие записи", font=MEDIUM_FONT).grid(row=3, column=0, pady=10, sticky="w", padx=12)

    # Таблица с записями
    tree_frame = ttk.Frame(page)
    tree_frame.grid(row=4, column=0, sticky="nsew", padx=12, pady=(0, 12))
    tree_frame.grid_rowconfigure(0, weight=1)
    tree_frame.grid_columnconfigure(0, weight=1)

    columns = ("time", "master", "service", "client")
    global main_tree
    main_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=8)
    main_tree.grid(row=0, column=0, sticky="nsew")

    # Настройка заголовков
    for col in columns:
        main_tree.heading(col, text=col.capitalize())
        main_tree.column(col, width=150)

    # Scrollbar
    scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=main_tree.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    main_tree.configure(yscrollcommand=scrollbar.set)

    # Заполняем данными
    refresh_main_page()


# ========== СТРАНИЦА ЗАПИСИ ==========
def refresh_booking_page():
    """Обновление страницы записи"""
    global selected_time_slot
    selected_time_slot = None

    # Обновляем список мастеров
    masters = get_masters()
    master_names = [f"{m['surname']} {m['name']}" for m in masters]
    master_combo['values'] = master_names
    if master_names:
        master_combo.set(master_names[0])
        update_master_schedule()

    # Обновляем список клиентов
    clients = get_clients()
    client_names = [c['full_name'] for c in clients]
    client_combo['values'] = client_names

    # Обновляем доступное время
    update_available_time()


def update_master_schedule():
    """Обновление расписания мастера"""
    global selected_master_id

    # Получаем ID выбранного мастера
    selected_master = master_combo.get()
    if selected_master:
        masters = get_masters()
        for master in masters:
            master_name = f"{master['surname']} {master['name']}"
            if master_name == selected_master:
                selected_master_id = master['id']
                break

    # Обновляем список услуг для выбранного мастера
    if selected_master_id:
        services = get_services_by_master(selected_master_id)
        service_names = [s['name'] for s in services]
        service_combo['values'] = service_names
        if service_names:
            service_combo.set(service_names[0])
            update_service_price()

    # Обновляем доступные даты
    update_date_restrictions()


def update_date_restrictions():
    """Обновление ограничений по дате"""
    if not selected_master_id:
        return

    # Получаем даты, когда работает мастер
    working_dates = get_master_working_dates(selected_master_id)

    # Настраиваем календарь
    def date_valid(date_str):
        if not working_dates:
            return False
        selected_date = datetime.strptime(date_str, '%d.%m.%Y').date()
        return selected_date in working_dates

    # Пока просто показываем сообщение
    if working_dates:
        dates_str = ", ".join([d.strftime('%d.%m') for d in working_dates[:5]])
        if len(working_dates) > 5:
            dates_str += f" ... (всего {len(working_dates)} дней)"
        date_label.config(text=f"Мастер работает: {dates_str}")
    else:
        date_label.config(text="У мастера нет установленного графика")


def update_available_time():
    """Обновление доступного времени"""
    if not selected_master_id:
        return

    # Получаем выбранную дату
    try:
        selected_date_str = cal.get_date()
        selected_date = datetime.strptime(selected_date_str, '%d.%m.%Y').date()
    except:
        return

    # Очищаем текущие слоты
    for widget in slots_frame.winfo_children():
        widget.destroy()

    # Получаем доступные слоты
    time_slots = get_available_time_slots(selected_master_id, selected_date)

    if not time_slots:
        no_slots_label = ttk.Label(slots_frame, text="Нет доступного времени", font=MEDIUM_FONT)
        no_slots_label.pack(pady=20)
        return

    # Отображаем слоты
    global time_slot_buttons
    time_slot_buttons = []

    for i, slot in enumerate(time_slots):
        row_frame = ttk.Frame(slots_frame, padding=(6, 4))
        row_frame.pack(fill="x", pady=4)
        row_frame.columnconfigure(0, weight=1)
        row_frame.columnconfigure(1, weight=1)
        row_frame.columnconfigure(2, weight=1)

        lbl_time = ttk.Label(row_frame, text=f"{slot['start']} - {slot['end']}", font=BIG_TIME_FONT, anchor="center")
        lbl_time.grid(row=0, column=0, sticky="nsew", padx=8)

        lbl_status = ttk.Label(row_frame, text="Свободно", font=MEDIUM_FONT, anchor="center", foreground="green")
        lbl_status.grid(row=0, column=1, sticky="nsew", padx=8)

        btn = ttk.Button(row_frame, text="Записать", style="Large.TButton",
                         command=lambda s=slot: select_time_slot(s))
        btn.grid(row=0, column=2, sticky="nsew", padx=8)
        time_slot_buttons.append(btn)

    # Обновляем canvas
    slots_frame.update_idletasks()
    canvas.configure(scrollregion=canvas.bbox("all"))


def update_service_price():
    """Обновление цены услуги"""
    selected_service = service_combo.get()
    if selected_service and selected_master_id:
        services = get_services_by_master(selected_master_id)
        for service in services:
            if service['name'] == selected_service:
                price_label.config(text=f"Цена: {service['cost']} р")
                return
    price_label.config(text="Цена: ____ р")


def select_time_slot(slot):
    """Выбор временного слота"""
    global selected_time_slot

    # Сбрасываем выделение предыдущего слота
    if selected_time_slot:
        selected_time_slot['button'].configure(style="Large.TButton")

    # Находим кнопку, соответствующую слоту
    for i, btn in enumerate(time_slot_buttons):
        if btn['command'].__defaults__[0] == slot:
            btn.configure(style="Selected.Menu.TButton")
            selected_time_slot = {
                'slot': slot,
                'button': btn
            }
            break


def on_client_selected(event):
    """Обработка выбора клиента"""
    selected_client = client_combo.get()
    if selected_client:
        clients = get_clients()
        for client in clients:
            if client['full_name'] == selected_client:
                phone_entry.delete(0, tk.END)
                phone_entry.insert(0, client['phone'])
                break


def on_service_selected(event):
    """Обработка выбора услуги"""
    update_service_price()


def confirm_booking():
    """Подтверждение записи"""
    global selected_time_slot

    if not selected_time_slot:
        messagebox.showerror("Ошибка", "Выберите время записи!")
        return

    # Получаем данные клиента
    client_name = client_combo.get()
    client_phone = phone_entry.get()

    if not client_name:
        messagebox.showerror("Ошибка", "Выберите клиента!")
        return

    # Находим ID клиента
    clients = get_clients()
    client_id = None
    for client in clients:
        if client['full_name'] == client_name:
            client_id = client['id']
            break

    if not client_id:
        messagebox.showerror("Ошибка", "Клиент не найден!")
        return

    # Получаем ID услуги
    service_name = service_combo.get()
    services = get_services_by_master(selected_master_id)
    service_id = None
    for service in services:
        if service['name'] == service_name:
            service_id = service['id']
            break

    if not service_id:
        messagebox.showerror("Ошибка", "Услуга не найдена!")
        return

    # Получаем дату
    try:
        selected_date_str = cal.get_date()
        selected_date = datetime.strptime(selected_date_str, '%d.%m.%Y').date()
    except:
        messagebox.showerror("Ошибка", "Некорректная дата!")
        return

    # Добавляем запись в БД
    try:
        add_appointment(
            client_id=client_id,
            master_id=selected_master_id,
            service_id=service_id,
            appointment_date=selected_date,
            start_time=selected_time_slot['slot']['start'],
            end_time=selected_time_slot['slot']['end']
        )

        messagebox.showinfo("Успех", "Запись успешно создана!")

        # Сбрасываем форму
        selected_time_slot = None
        update_available_time()

        # Переходим на главную страницу
        show_page("Главная")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось создать запись: {str(e)}")


def create_booking_page(page):
    """Создание страницы записи"""
    global selected_time_slot, selected_master_id
    selected_time_slot = None
    selected_master_id = None

    page.grid_rowconfigure(4, weight=0)
    page.grid_rowconfigure(3, weight=0)
    page.grid_rowconfigure(2, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="Запись", font=LARGE_FONT).grid(row=0, column=0, pady=10, sticky="n")

    # Верхняя панель с выбором даты и мастера
    top_frame = ttk.Frame(page)
    top_frame.grid(row=1, column=0, sticky="ew", pady=(0, 10), padx=12)
    top_frame.grid_columnconfigure(0, weight=1)

    # Дата
    date_frame = ttk.Frame(top_frame)
    date_frame.grid(row=0, column=0, sticky="w", pady=4)
    ttk.Label(date_frame, text="Дата:", font=MEDIUM_FONT).pack(side="left")

    global cal
    cal = DateEntry(date_frame, width=14, date_pattern='dd.mm.yyyy')
    cal.pack(side="left", padx=6)
    make_large_entry(cal)

    # Обработчик изменения даты
    def on_date_selected(event):
        update_available_time()

    cal.bind("<<DateEntrySelected>>", on_date_selected)

    # Мастер
    master_frame = ttk.Frame(top_frame)
    master_frame.grid(row=1, column=0, sticky="w", pady=4)
    ttk.Label(master_frame, text="Мастер:", font=MEDIUM_FONT).pack(side="left")

    global master_combo
    master_combo = ttk.Combobox(master_frame, state="readonly")
    master_combo.pack(side="left", padx=6)
    master_combo.bind("<<ComboboxSelected>>", lambda e: update_master_schedule())
    make_large_combo(master_combo)

    # Метка с информацией о доступных датах
    global date_label
    date_label = ttk.Label(top_frame, text="", font=SMALL_FONT, foreground="blue")
    date_label.grid(row=2, column=0, sticky="w", pady=4)

    # Центральная область с доступным временем
    center_container = ttk.Frame(page)
    center_container.grid(row=2, column=0, sticky="nsew", pady=6, padx=12)
    center_container.grid_rowconfigure(0, weight=1)
    center_container.grid_columnconfigure(0, weight=1)

    boxed = ttk.Frame(center_container, relief="groove", borderwidth=1)
    boxed.grid(row=0, column=0, sticky="nsew")
    boxed.grid_rowconfigure(0, weight=1)
    boxed.grid_columnconfigure(0, weight=1)

    global canvas, slots_frame
    canvas = tk.Canvas(boxed, highlightthickness=0)
    scrollbar = ttk.Scrollbar(boxed, orient="vertical", command=canvas.yview)
    canvas.configure(yscrollcommand=scrollbar.set)

    scrollbar.grid(row=0, column=1, sticky="ns")
    canvas.grid(row=0, column=0, sticky="nsew")

    slots_frame = ttk.Frame(canvas)
    canvas.create_window((0, 0), window=slots_frame, anchor="nw")

    def on_frame_configure(event):
        canvas.configure(scrollregion=canvas.bbox("all"))

    slots_frame.bind("<Configure>", on_frame_configure)

    def _on_mousewheel(event):
        try:
            if sys.platform == "darwin":
                delta = int(-1 * event.delta)
            else:
                delta = int(-1 * (event.delta / 120))
        except Exception:
            delta = -1 if getattr(event, "delta", 0) > 0 else 1
        canvas.yview_scroll(delta, "units")

    canvas.bind("<MouseWheel>", _on_mousewheel)
    canvas.bind("<Button-4>", lambda e: canvas.yview_scroll(-1, "units"))
    canvas.bind("<Button-5>", lambda e: canvas.yview_scroll(1, "units"))

    canvas.config(height=360)

    # Информация о клиенте и услуге
    info_frame = ttk.Frame(page, relief="groove", borderwidth=1, padding=10)
    info_frame.grid(row=3, column=0, sticky="ew", pady=(10, 0), padx=12)
    info_frame.grid_columnconfigure(0, weight=0)
    info_frame.grid_columnconfigure(1, weight=1)
    info_frame.grid_columnconfigure(2, weight=0)
    info_frame.grid_columnconfigure(3, weight=1)
    info_frame.grid_columnconfigure(4, weight=0)

    # Клиент
    ttk.Label(info_frame, text="Клиент:", font=MEDIUM_FONT).grid(row=0, column=0, sticky="w", pady=4, padx=(0, 4))

    global client_combo
    client_combo = ttk.Combobox(info_frame, state="readonly")
    client_combo.grid(row=0, column=1, sticky="ew", pady=4, padx=(0, 12))
    client_combo.bind("<<ComboboxSelected>>", on_client_selected)
    make_large_combo(client_combo)

    # Телефон
    ttk.Label(info_frame, text="Телефон:", font=MEDIUM_FONT).grid(row=0, column=2, sticky="w", pady=4, padx=(0, 4))

    global phone_entry
    phone_entry = ttk.Entry(info_frame)
    phone_entry.grid(row=0, column=3, sticky="ew", pady=4, padx=(0, 12))
    make_large_entry(phone_entry)

    # Кнопка "Новый клиент"
    ttk.Button(info_frame, text="Новый клиент", style="Large.TButton",
               command=lambda: show_page("Клиенты")).grid(row=0, column=4, pady=4)

    # Услуга
    ttk.Label(info_frame, text="Услуга:", font=MEDIUM_FONT).grid(row=1, column=0, sticky="w", pady=8, padx=(0, 4))

    global service_combo
    service_combo = ttk.Combobox(info_frame, state="readonly")
    service_combo.grid(row=1, column=1, sticky="ew", pady=8, padx=(0, 12))
    service_combo.bind("<<ComboboxSelected>>", on_service_selected)
    make_large_combo(service_combo)

    # Цена
    global price_label
    price_label = ttk.Label(info_frame, text="Цена: ____ р", font=MEDIUM_FONT)
    price_label.grid(row=1, column=2, columnspan=2, sticky="w", pady=8)

    # Кнопки действий
    action_frame = ttk.Frame(page)
    action_frame.grid(row=4, column=0, pady=(20, 10), padx=12, sticky="ew")
    action_frame.grid_columnconfigure(0, weight=1)
    action_frame.grid_columnconfigure(1, weight=1)
    action_frame.grid_columnconfigure(2, weight=1)

    ttk.Button(action_frame, text="Подтвердить запись", style="Large.TButton",
               command=confirm_booking).grid(row=0, column=0, padx=8, sticky="ew")
    ttk.Button(action_frame, text="Изменить", style="Large.TButton").grid(row=0, column=1, padx=8, sticky="ew")
    ttk.Button(action_frame, text="Удалить", style="Large.TButton").grid(row=0, column=2, padx=8, sticky="ew")


# ========== СТРАНИЦА ГРАФИКА РАБОТЫ ==========
def refresh_schedule_page():
    """Обновление страницы графика работы"""
    # Обновляем список мастеров
    masters = get_masters()
    master_names = [f"{m['surname']} {m['name']}" for m in masters]
    master_combo_sched['values'] = master_names
    if master_names:
        master_combo_sched.set(master_names[0])
        load_master_schedule()





def load_master_schedule():
    """Загрузка расписания мастера"""
    try:
        # Получаем ID выбранного мастера
        selected_master = master_combo_sched.get()
        if not selected_master:
            return

        masters = get_masters()
        master_id = None
        for master in masters:
            master_name = f"{master['surname']} {master['name']}"
            if master_name == selected_master:
                master_id = master['id']
                break

        if not master_id:
            return

        # Получаем выбранную дату из календаря
        try:
            # Получаем дату из DateEntry
            selected_date_str = week_cal.get()

            # Преобразуем строку в дату
            # DateEntry возвращает строку в формате 'dd.mm.yyyy'
            week_start = datetime.strptime(selected_date_str, '%d.%m.%Y').date()

            # Находим понедельник выбранной недели
            # Понедельник = 0, Воскресенье = 6 в Python
            while week_start.weekday() != 0:  # 0 = понедельник
                week_start -= timedelta(days=1)

        except Exception as e:
            print(f"Ошибка парсинга даты из календаря: {e}")
            # Если ошибка, используем текущую неделю
            week_start = date.today()
            while week_start.weekday() != 0:
                week_start -= timedelta(days=1)

            # Устанавливаем дату в календарь
            week_cal.set_date(week_start)

        # Вычисляем конец недели (воскресенье)
        week_end = week_start + timedelta(days=6)

        # Обновляем метку с информацией о неделе
        if 'week_info_label' in globals():
            week_info_label.config(
                text=f"Неделя {week_start.strftime('%d.%m.%Y')} - {week_end.strftime('%d.%m.%Y')}"
            )

        # Загружаем расписание на всю неделю
        try:
            schedule = get_master_schedule(master_id, week_start)
        except Exception as e:
            print(f"Ошибка загрузки расписания: {e}")
            schedule = []

        # Обновляем данные в таблице для всех 7 дней недели
        for i in range(7):
            current_date = week_start + timedelta(days=i)

            # Обновляем метку с датой
            if i < len(day_labels):
                day_labels[i].config(
                    text=f"{current_date.strftime('%d.%m')} ({['Пн', 'Вт', 'Ср', 'Чт', 'Пт', 'Сб', 'Вс'][i]})"
                )

            # Ищем запись на эту дату
            day_schedule = None
            for s in schedule:
                if s['work_date'] == current_date:
                    day_schedule = s
                    break

            # Устанавливаем значения
            if i < len(day_vars) and i < len(start_combos) and i < len(end_combos):
                if day_schedule:
                    day_vars[i].set(1)
                    # Преобразуем время в строку
                    start_time = day_schedule['start_time']
                    end_time = day_schedule['end_time']

                    if isinstance(start_time, str):
                        start_time_str = start_time[:5] if ':' in start_time else "09:00"
                    else:
                        start_time_str = start_time.strftime('%H:%M') if start_time else "09:00"

                    if isinstance(end_time, str):
                        end_time_str = end_time[:5] if ':' in end_time else "18:00"
                    else:
                        end_time_str = end_time.strftime('%H:%M') if end_time else "18:00"

                    start_combos[i].set(start_time_str)
                    end_combos[i].set(end_time_str)
                else:
                    # По умолчанию: будни - работает, выходные - нет
                    day_vars[i].set(0 if current_date.weekday() >= 5 else 1)
                    start_combos[i].set("09:00")
                    end_combos[i].set("18:00")

    except Exception as e:
        print(f"Общая ошибка в load_master_schedule: {e}")
        messagebox.showerror("Ошибка", f"Не удалось загрузить расписание: {str(e)}")


def save_schedule():
    """Сохранение расписания мастера"""
    try:
        # Получаем ID мастера
        selected_master = master_combo_sched.get()
        if not selected_master:
            messagebox.showerror("Ошибка", "Выберите мастера!")
            return

        masters = get_masters()
        master_id = None
        for master in masters:
            master_name = f"{master['surname']} {master['name']}"
            if master_name == selected_master:
                master_id = master['id']
                break

        if not master_id:
            messagebox.showerror("Ошибка", "Мастер не найден!")
            return

        # Получаем выбранную неделю из календаря
        try:
            # Получаем дату из DateEntry
            selected_date_str = week_cal.get()

            # Преобразуем строку в дату
            week_start = datetime.strptime(selected_date_str, '%d.%m.%Y').date()

            # Находим понедельник
            while week_start.weekday() != 0:
                week_start -= timedelta(days=1)

        except Exception as e:
            messagebox.showerror("Ошибка",
                                 f"Некорректная дата в календаре! Формат должен быть дд.мм.гггг\nОшибка: {str(e)}")
            return

        # Собираем данные для всех дней недели
        day_data = []
        errors = []

        for i in range(7):
            current_date = week_start + timedelta(days=i)

            if day_vars[i].get() == 1:  # Если мастер работает в этот день
                # Проверяем время
                start_time = start_combos[i].get()
                end_time = end_combos[i].get()

                if not start_time or not end_time:
                    errors.append(f"Укажите время работы для {current_date.strftime('%d.%m.%Y')}")
                    continue

                try:
                    # Проверяем формат времени
                    start_dt = datetime.strptime(start_time, '%H:%M')
                    end_dt = datetime.strptime(end_time, '%H:%M')

                    # Проверяем, что начало раньше конца
                    if start_dt >= end_dt:
                        errors.append(
                            f"Время начала должно быть раньше времени окончания для {current_date.strftime('%d.%m.%Y')}")
                        continue

                except ValueError:
                    errors.append(
                        f"Некорректное время для {current_date.strftime('%d.%m.%Y')}. Формат: ЧЧ:ММ (например, 09:00)")
                    continue

                day_data.append({
                    'day_of_week': i + 1,  # 1=Пн, 7=Вс
                    'work_date': current_date,
                    'start_time': start_time,
                    'end_time': end_time
                })
            else:
                # Если мастер не работает в этот день, все равно добавляем запись
                # с пустым временем или не добавляем вообще
                # Это зависит от логики вашего приложения
                pass

        # Если есть ошибки, показываем их
        if errors:
            error_text = "\n".join(errors[:5])  # Показываем первые 5 ошибок
            if len(errors) > 5:
                error_text += f"\n\n... и еще {len(errors) - 5} ошибок"
            messagebox.showerror("Ошибки в данных", error_text)
            return

        # Сохраняем в БД
        try:
            result = set_master_schedule(master_id, day_data)
            if result >= 0:
                messagebox.showinfo("Успех", f"Расписание на неделю сохранено! Сохранено {result} дней.")

                # Обновляем страницу записи, если она открыта
                if current_page == "Запись":
                    refresh_booking_page()

                # Перезагружаем расписание для отображения изменений
                load_master_schedule()
            else:
                messagebox.showerror("Ошибка", "Не удалось сохранить расписание!")

        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить расписание в базу данных: {str(e)}")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Неожиданная ошибка: {str(e)}")







def create_schedule_page(page):
    """Создание страницы графика работы"""
    page.grid_rowconfigure(1, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="График работы", font=LARGE_FONT).grid(row=0, column=0, pady=12, sticky="n")

    main_frame = ttk.Frame(page)
    main_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=6)
    main_frame.grid_columnconfigure(0, weight=1)

    # Верхняя панель
    top_frame = ttk.Frame(main_frame)
    top_frame.grid(row=0, column=0, sticky="ew", pady=6)

    ttk.Label(top_frame, text="Мастер:", font=MEDIUM_FONT).grid(row=0, column=0, sticky="w")

    global master_combo_sched
    master_combo_sched = ttk.Combobox(top_frame, state="readonly")
    master_combo_sched.grid(row=0, column=1, sticky="w", padx=8)
    master_combo_sched.bind("<<ComboboxSelected>>", lambda e: load_master_schedule())
    make_large_combo(master_combo_sched)

    ttk.Label(top_frame, text="Неделя с:", font=MEDIUM_FONT).grid(row=0, column=2, sticky="w", padx=(12, 0))

    global week_cal
    week_cal = DateEntry(top_frame, width=14, date_pattern='dd.mm.yyyy')
    week_cal.grid(row=0, column=3, sticky="w", padx=8)
    week_cal.bind("<<DateEntrySelected>>", lambda e: load_master_schedule())
    make_large_entry(week_cal)

    # Кнопки навигации по неделям
    nav_frame = ttk.Frame(top_frame)
    nav_frame.grid(row=0, column=4, sticky="w", padx=(12, 0))

    def prev_week():
        try:
            current_date_str = week_cal.get_date()
            current_date = datetime.strptime(current_date_str, '%d.%m.%Y').date()
            new_date = current_date - timedelta(days=7)
            week_cal.set_date(new_date.strftime('%d.%m.%Y'))
            load_master_schedule()
        except:
            pass

    def next_week():
        try:
            current_date_str = week_cal.get_date()
            current_date = datetime.strptime(current_date_str, '%d.%m.%Y').date()
            new_date = current_date + timedelta(days=7)
            week_cal.set_date(new_date.strftime('%d.%m.%Y'))
            load_master_schedule()
        except:
            pass

    ttk.Button(nav_frame, text="◀ Пред.", command=prev_week).pack(side="left", padx=2)
    ttk.Button(nav_frame, text="След. ▶", command=next_week).pack(side="left", padx=2)

    # Метка с информацией о неделе
    global week_info_label
    week_info_label = ttk.Label(top_frame, text="", font=SMALL_FONT, foreground="blue")
    week_info_label.grid(row=1, column=0, columnspan=5, sticky="w", pady=4)

    # Таблица расписания
    schedule_frame = ttk.Frame(main_frame, relief="groove", borderwidth=1, padding=10)
    schedule_frame.grid(row=1, column=0, sticky="ew", pady=6)

    for c in range(1, 8):
        schedule_frame.grid_columnconfigure(c, weight=1, uniform="schedcol")

    # Создаем глобальные переменные для хранения данных
    global day_vars, start_combos, end_combos, day_labels
    day_vars = []
    start_combos = []
    end_combos = []
    day_labels = []

    time_options = [f"{h:02d}:00" for h in range(8, 23)]

    # Заголовки столбцов
    days_of_week = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
    ttk.Label(schedule_frame, text="", font=SMALL_FONT).grid(row=0, column=0)
    for i, day in enumerate(days_of_week):
        day_label = ttk.Label(schedule_frame, text=day, font=MEDIUM_FONT)
        day_label.grid(row=0, column=i + 1, sticky="nsew", padx=4, pady=6)
        day_labels.append(day_label)

    # Работает
    ttk.Label(schedule_frame, text="Работает", font=MEDIUM_FONT).grid(row=1, column=0, sticky="w", padx=6)
    for i in range(7):
        var = tk.IntVar(value=1 if i < 5 else 0)
        day_vars.append(var)
        chk = tk.Checkbutton(schedule_frame, variable=var, onvalue=1, offvalue=0,
                             font=MEDIUM_FONT, text="", padx=6, pady=4)
        chk.grid(row=1, column=i + 1, sticky="nsew", padx=4, pady=6)

    # Начало работы
    ttk.Label(schedule_frame, text="Начало", font=MEDIUM_FONT).grid(row=2, column=0, sticky="w", padx=6, pady=6)
    for i in range(7):
        combo = ttk.Combobox(schedule_frame, values=time_options, state="readonly", width=10)
        combo.set("09:00")
        combo.grid(row=2, column=i + 1, sticky="nsew", padx=4, pady=6)
        start_combos.append(combo)
        make_large_combo(combo)

    # Конец работы
    ttk.Label(schedule_frame, text="Конец", font=MEDIUM_FONT).grid(row=3, column=0, sticky="w", padx=6, pady=6)
    for i in range(7):
        combo = ttk.Combobox(schedule_frame, values=time_options, state="readonly", width=10)
        combo.set("18:00")
        combo.grid(row=3, column=i + 1, sticky="nsew", padx=4, pady=6)
        end_combos.append(combo)
        make_large_combo(combo)

    # Кнопка сохранения
    btn_frame = ttk.Frame(main_frame)
    btn_frame.grid(row=2, column=0, sticky="w", pady=6)
    ttk.Button(btn_frame, text="Сохранить расписание", style="Large.TButton",
               command=save_schedule).pack(side="left", padx=4)

    # Загружаем данные при открытии страницы
    refresh_schedule_page()


# def save_schedule():
#     """Сохранение расписания мастера"""
#     # Получаем ID мастера
#     selected_master = master_combo_sched.get()
#     if not selected_master:
#         messagebox.showerror("Ошибка", "Выберите мастера!")
#         return
#
#     masters = get_masters()
#     master_id = None
#     for master in masters:
#         master_name = f"{master['surname']} {master['name']}"
#         if master_name == selected_master:
#             master_id = master['id']
#             break
#
#     if not master_id:
#         messagebox.showerror("Ошибка", "Мастер не найден!")
#         return
#
#     # Получаем выбранную неделю
#     try:
#         selected_date = week_cal.get_date()
#         week_start = datetime.strptime(selected_date, '%d.%m.%Y').date()
#
#         # Находим понедельник
#         while week_start.weekday() != 0:
#             week_start -= timedelta(days=1)
#     except:
#         messagebox.showerror("Ошибка", "Некорректная дата!")
#         return
#
#     # Собираем данные
#     day_data = []
#     for i in range(7):
#         if day_vars[i].get() == 1:  # Если мастер работает в этот день
#             current_date = week_start + timedelta(days=i)
#
#             # Проверяем время
#             start_time = start_combos[i].get()
#             end_time = end_combos[i].get()
#
#             if not start_time or not end_time:
#                 messagebox.showerror("Ошибка", f"Укажите время работы для {current_date.strftime('%d.%m')}")
#                 return
#
#             try:
#                 # Проверяем формат времени
#                 datetime.strptime(start_time, '%H:%M')
#                 datetime.strptime(end_time, '%H:%M')
#             except:
#                 messagebox.showerror("Ошибка", f"Некорректное время для {current_date.strftime('%d.%m')}")
#                 return
#
#             day_data.append({
#                 'day_of_week': i + 1,
#                 'work_date': current_date,
#                 'start_time': start_time,
#                 'end_time': end_time
#             })
#
#     # Сохраняем в БД
#     try:
#         set_master_schedule(master_id, day_data)
#         messagebox.showinfo("Успех", "Расписание успешно сохранено!")
#
#         # Обновляем страницу записи, если она открыта
#         if current_page == "Запись":
#             refresh_booking_page()
#
#     except Exception as e:
#         messagebox.showerror("Ошибка", f"Не удалось сохранить расписание: {str(e)}")


# def create_schedule_page(page):
#     """Создание страницы графика работы"""
#     page.grid_rowconfigure(1, weight=1)
#     page.grid_columnconfigure(0, weight=1)
#
#     # Заголовок
#     ttk.Label(page, text="График работы", font=LARGE_FONT).grid(row=0, column=0, pady=12, sticky="n")
#
#     main_frame = ttk.Frame(page)
#     main_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=6)
#     main_frame.grid_columnconfigure(0, weight=1)
#
#     # Верхняя панель
#     top_frame = ttk.Frame(main_frame)
#     top_frame.grid(row=0, column=0, sticky="ew", pady=6)
#
#     ttk.Label(top_frame, text="Мастер:", font=MEDIUM_FONT).grid(row=0, column=0, sticky="w")
#
#     global master_combo_sched
#     master_combo_sched = ttk.Combobox(top_frame, state="readonly")
#     master_combo_sched.grid(row=0, column=1, sticky="w", padx=8)
#     master_combo_sched.bind("<<ComboboxSelected>>", lambda e: load_master_schedule())
#     make_large_combo(master_combo_sched)
#
#     ttk.Label(top_frame, text="Неделя с:", font=MEDIUM_FONT).grid(row=0, column=2, sticky="w", padx=(12, 0))
#
#     global week_cal
#     week_cal = DateEntry(top_frame, width=14, date_pattern='dd.mm.yyyy')
#     week_cal.grid(row=0, column=3, sticky="w", padx=8)
#     week_cal.bind("<<DateEntrySelected>>", lambda e: load_master_schedule())
#     make_large_entry(week_cal)
#
#     # Таблица расписания
#     schedule_frame = ttk.Frame(main_frame, relief="groove", borderwidth=1, padding=10)
#     schedule_frame.grid(row=1, column=0, sticky="ew", pady=6)
#
#     for c in range(1, 8):
#         schedule_frame.grid_columnconfigure(c, weight=1, uniform="schedcol")
#
#     # Создаем глобальные переменные для хранения данных
#     global day_vars, start_combos, end_combos, day_labels
#     day_vars = []
#     start_combos = []
#     end_combos = []
#     day_labels = []
#
#     time_options = [f"{h:02d}:00" for h in range(8, 23)]
#
#     # Заголовки столбцов
#     days_of_week = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]
#     ttk.Label(schedule_frame, text="", font=SMALL_FONT).grid(row=0, column=0)
#     for i, day in enumerate(days_of_week):
#         day_label = ttk.Label(schedule_frame, text=day, font=MEDIUM_FONT)
#         day_label.grid(row=0, column=i + 1, sticky="nsew", padx=4, pady=6)
#         day_labels.append(day_label)
#
#     # Работает
#     ttk.Label(schedule_frame, text="Работает", font=MEDIUM_FONT).grid(row=1, column=0, sticky="w", padx=6)
#     for i in range(7):
#         var = tk.IntVar(value=1 if i < 5 else 0)
#         day_vars.append(var)
#         chk = tk.Checkbutton(schedule_frame, variable=var, onvalue=1, offvalue=0,
#                              font=MEDIUM_FONT, text="", padx=6, pady=4)
#         chk.grid(row=1, column=i + 1, sticky="nsew", padx=4, pady=6)
#
#     # Начало работы
#     ttk.Label(schedule_frame, text="Начало", font=MEDIUM_FONT).grid(row=2, column=0, sticky="w", padx=6, pady=6)
#     for i in range(7):
#         combo = ttk.Combobox(schedule_frame, values=time_options, state="readonly", width=10)
#         combo.set("09:00")
#         combo.grid(row=2, column=i + 1, sticky="nsew", padx=4, pady=6)
#         start_combos.append(combo)
#         make_large_combo(combo)
#
#     # Конец работы
#     ttk.Label(schedule_frame, text="Конец", font=MEDIUM_FONT).grid(row=3, column=0, sticky="w", padx=6, pady=6)
#     for i in range(7):
#         combo = ttk.Combobox(schedule_frame, values=time_options, state="readonly", width=10)
#         combo.set("18:00")
#         combo.grid(row=3, column=i + 1, sticky="nsew", padx=4, pady=6)
#         end_combos.append(combo)
#         make_large_combo(combo)
#
#     # Кнопка сохранения
#     btn_frame = ttk.Frame(main_frame)
#     btn_frame.grid(row=2, column=0, sticky="w", pady=6)
#     ttk.Button(btn_frame, text="Сохранить расписание", style="Large.TButton",
#                command=save_schedule).pack(side="left", padx=4)


# ========== СТРАНИЦА УСЛУГ ==========
def refresh_services_page():
    """Обновление страницы услуг"""
    # Очищаем таблицу
    for item in services_tree.get_children():
        services_tree.delete(item)

    # Загружаем услуги из БД
    services = get_services()
    for service in services:
        services_tree.insert("", tk.END, values=(
            service['id'],  # Скрытый ID
            service['name'],
            f"{service['duration']} мин",
            f"{service['cost']} р",
            service['master_name']
        ))

    # Скрываем первый столбец с ID
    services_tree.column("#1", width=0, stretch=False)
    services_tree.heading("#1", text="")


def on_service_select(event):
    """Обработка выбора услуги в таблице"""
    selection = services_tree.selection()
    if selection:
        item = services_tree.item(selection[0])
        values = item['values']

        # Заполняем поля формы данными выбранной услуги
        if len(values) >= 5:
            # ID в values[0] (скрытый)
            service_name_entry.delete(0, tk.END)
            service_name_entry.insert(0, values[1])

            # Извлекаем число из строки "60 мин"
            duration_str = values[2].replace(' мин', '').strip()
            service_duration_entry.delete(0, tk.END)
            service_duration_entry.insert(0, duration_str)

            # Извлекаем число из строки "1000 р"
            cost_str = values[3].replace(' р', '').strip()
            service_cost_entry.delete(0, tk.END)
            service_cost_entry.insert(0, cost_str)

            # Устанавливаем мастера
            service_master_combo.set(values[4])


def edit_service():
    """Редактирование выбранной услуги"""
    selection = services_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите услугу для редактирования!")
        return

    item = services_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 5:
        messagebox.showerror("Ошибка", "Не удалось получить данные услуги!")
        return

    service_id = values[0]  # Скрытый ID

    # Получаем данные из формы
    name = service_name_entry.get()
    duration = service_duration_entry.get()
    cost = service_cost_entry.get()
    master_name = service_master_combo.get()

    if not all([name, duration, cost, master_name]):
        messagebox.showerror("Ошибка", "Заполните все поля!")
        return

    # Находим ID мастера
    masters = get_masters()
    master_id = None
    for master in masters:
        master_full_name = f"{master['surname']} {master['name']}"
        if master_full_name == master_name:
            master_id = master['id']
            break

    if not master_id:
        messagebox.showerror("Ошибка", "Мастер не найден!")
        return

    # Проверяем числовые значения
    try:
        duration_int = int(duration)
        cost_int = int(cost)
    except ValueError:
        messagebox.showerror("Ошибка", "Длительность и стоимость должны быть числами!")
        return

    # Обновляем услугу
    try:
        result = update_service(service_id, name, duration_int, cost_int, master_id)
        if result > 0:
            messagebox.showinfo("Успех", "Услуга обновлена!")
            refresh_services_page()

            # Очищаем форму
            service_name_entry.delete(0, tk.END)
            service_duration_entry.delete(0, tk.END)
            service_cost_entry.delete(0, tk.END)
            service_master_combo.set('')
        else:
            messagebox.showerror("Ошибка", "Не удалось обновить услугу!")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось обновить услугу: {str(e)}")


def delete_service_from_db():
    """Удаление выбранной услуги"""
    selection = services_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите услугу для удаления!")
        return

    item = services_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 5:
        messagebox.showerror("Ошибка", "Не удалось получить данные услуги!")
        return

    service_id = values[0]  # Скрытый ID

    # Подтверждение удаления
    if not messagebox.askyesno("Подтверждение", "Вы уверены, что хотите удалить эту услугу?"):
        return

    # Удаляем услугу
    try:
        result = delete_service(service_id)
        if result > 0:
            messagebox.showinfo("Успех", "Услуга удалена!")
            refresh_services_page()

            # Очищаем форму
            service_name_entry.delete(0, tk.END)
            service_duration_entry.delete(0, tk.END)
            service_cost_entry.delete(0, tk.END)
            service_master_combo.set('')
        else:
            messagebox.showerror("Ошибка", "Не удалось удалить услугу!")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось удалить услугу: {str(e)}")

    # Загружаем услуги из БД
    services = get_services()
    for service in services:
        services_tree.insert("", tk.END, values=(
            service['name'],
            f"{service['duration']} мин",
            f"{service['cost']} р",
            service['master_name']
        ))


def add_service_from_form():
    """Добавление услуги из формы"""
    name = service_name_entry.get()
    duration = service_duration_entry.get()
    cost = service_cost_entry.get()
    master_name = service_master_combo.get()

    if not all([name, duration, cost, master_name]):
        messagebox.showerror("Ошибка", "Заполните все поля!")
        return

    # Находим ID мастера
    masters = get_masters()
    master_id = None
    for master in masters:
        master_full_name = f"{master['surname']} {master['name']}"
        if master_full_name == master_name:
            master_id = master['id']
            break

    if not master_id:
        messagebox.showerror("Ошибка", "Мастер не найден!")
        return

    # Проверяем числовые значения
    try:
        duration_int = int(duration)
        cost_int = int(cost)
    except ValueError:
        messagebox.showerror("Ошибка", "Длительность и стоимость должны быть числами!")
        return

    # Добавляем услугу
    try:
        add_service(name, duration_int, cost_int, master_id)
        messagebox.showinfo("Успех", "Услуга добавлена!")

        # Очищаем форму
        service_name_entry.delete(0, tk.END)
        service_duration_entry.delete(0, tk.END)
        service_cost_entry.delete(0, tk.END)
        service_master_combo.set('')

        # Обновляем таблицу
        refresh_services_page()

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось добавить услугу: {str(e)}")


def create_services_page(page):
    """Создание страницы услуг"""
    page.grid_rowconfigure(2, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="Услуги", font=LARGE_FONT).grid(row=0, column=0, pady=10, sticky="n")

    # Форма добавления/редактирования
    add_frame = ttk.Frame(page, padding=8, relief="groove", borderwidth=1)
    add_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=6)
    add_frame.grid_columnconfigure(0, weight=0)
    add_frame.grid_columnconfigure(1, weight=1)

    ttk.Label(add_frame, text="Название услуги:", font=MEDIUM_FONT).grid(row=0, column=0, sticky="w", padx=4, pady=4)
    global service_name_entry
    service_name_entry = ttk.Entry(add_frame)
    service_name_entry.grid(row=0, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(service_name_entry)

    ttk.Label(add_frame, text="Длительность (мин):", font=MEDIUM_FONT).grid(row=1, column=0, sticky="w", padx=4, pady=4)
    global service_duration_entry
    service_duration_entry = ttk.Entry(add_frame)
    service_duration_entry.grid(row=1, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(service_duration_entry)

    ttk.Label(add_frame, text="Стоимость:", font=MEDIUM_FONT).grid(row=2, column=0, sticky="w", padx=4, pady=4)
    global service_cost_entry
    service_cost_entry = ttk.Entry(add_frame)
    service_cost_entry.grid(row=2, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(service_cost_entry)

    ttk.Label(add_frame, text="Мастер:", font=MEDIUM_FONT).grid(row=3, column=0, sticky="w", padx=4, pady=4)

    # Заполняем список мастеров
    masters = get_masters()
    master_names = [f"{m['surname']} {m['name']}" for m in masters]

    global service_master_combo
    service_master_combo = ttk.Combobox(add_frame, values=master_names, state="readonly")
    service_master_combo.grid(row=3, column=1, sticky="ew", padx=4, pady=4)
    make_large_combo(service_master_combo)

    # Кнопки формы
    form_buttons_frame = ttk.Frame(add_frame)
    form_buttons_frame.grid(row=4, column=0, columnspan=2, pady=6)

    ttk.Button(form_buttons_frame, text="Добавить", style="Large.TButton",
               command=add_service_from_form).pack(side="left", padx=4)
    ttk.Button(form_buttons_frame, text="Очистить", style="Large.TButton",
               command=lambda: [service_name_entry.delete(0, tk.END),
                                service_duration_entry.delete(0, tk.END),
                                service_cost_entry.delete(0, tk.END),
                                service_master_combo.set('')]).pack(side="left", padx=4)

    # Таблица услуг
    tree_frame = ttk.Frame(page)
    tree_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=(6, 0))
    tree_frame.grid_rowconfigure(0, weight=1)
    tree_frame.grid_columnconfigure(0, weight=1)

    # Добавляем столбец для ID (будет скрыт)
    columns = ("id", "name", "duration", "cost", "master")
    global services_tree
    services_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=15)
    services_tree.grid(row=0, column=0, sticky="nsew")

    # Настройка столбцов (первый скрыт)
    services_tree.heading("name", text="Название")
    services_tree.heading("duration", text="Длительность")
    services_tree.heading("cost", text="Стоимость")
    services_tree.heading("master", text="Мастер")

    services_tree.column("name", width=200)
    services_tree.column("duration", width=120)
    services_tree.column("cost", width=120)
    services_tree.column("master", width=150)

    # Привязываем обработчик выбора
    services_tree.bind("<<TreeviewSelect>>", on_service_select)

    # Scrollbar
    scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=services_tree.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    services_tree.configure(yscrollcommand=scrollbar.set)

    # Кнопки действий под таблицей
    action_frame = ttk.Frame(page)
    action_frame.grid(row=3, column=0, pady=10, padx=12, sticky="ew")
    action_frame.grid_columnconfigure(0, weight=1)
    action_frame.grid_columnconfigure(1, weight=1)

    ttk.Button(action_frame, text="Изменить", style="Large.TButton",
               command=edit_service).grid(row=0, column=0, padx=4, sticky="ew")
    ttk.Button(action_frame, text="Удалить", style="Large.TButton",
               command=delete_service_from_db).grid(row=0, column=1, padx=4, sticky="ew")

    # Загружаем данные
    refresh_services_page()


# ========== СТРАНИЦА КЛИЕНТОВ ==========
def refresh_clients_page():
    """Обновление страницы клиентов"""
    # Очищаем таблицу
    for item in clients_tree.get_children():
        clients_tree.delete(item)

    # Загружаем клиентов из БД
    clients = get_clients()
    for client in clients:
        clients_tree.insert("", tk.END, values=(
            client['id'],  # Скрытый ID
            client['full_name'],
            client['phone']
        ))

    # Скрываем первый столбец с ID
    clients_tree.column("#1", width=0, stretch=False)
    clients_tree.heading("#1", text="")


def on_client_select(event):
    """Обработка выбора клиента в таблице"""
    selection = clients_tree.selection()
    if selection:
        item = clients_tree.item(selection[0])
        values = item['values']

        # Заполняем поля формы данными выбранного клиента
        if len(values) >= 3:
            client_name_entry.delete(0, tk.END)
            client_name_entry.insert(0, values[1])

            client_phone_entry.delete(0, tk.END)
            client_phone_entry.insert(0, values[2])


def edit_client():
    """Редактирование выбранного клиента"""
    selection = clients_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите клиента для редактирования!")
        return

    item = clients_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 3:
        messagebox.showerror("Ошибка", "Не удалось получить данные клиента!")
        return

    client_id = values[0]  # Скрытый ID

    # Получаем данные из формы
    name = client_name_entry.get()
    phone = client_phone_entry.get()

    if not name or not phone:
        messagebox.showerror("Ошибка", "Заполните все поля!")
        return

    # Проверяем формат телефона
    if not validate_phone(phone):
        messagebox.showerror("Ошибка", "Неверный формат телефона! Должен быть: + и 11 цифр")
        return

    # Разбиваем имя на части
    name_parts = name.split()
    surname = name_parts[0] if len(name_parts) > 0 else ""
    first_name = name_parts[1] if len(name_parts) > 1 else name_parts[0] if name_parts else ""
    patronymic = name_parts[2] if len(name_parts) > 2 else ""

    # Обновляем клиента
    try:
        result = update_client(client_id, surname, first_name, patronymic, phone)
        if result > 0:
            messagebox.showinfo("Успех", "Клиент обновлен!")
            refresh_clients_page()

            # Очищаем форму
            client_name_entry.delete(0, tk.END)
            client_phone_entry.delete(0, tk.END)
        else:
            messagebox.showerror("Ошибка", "Не удалось обновить клиента!")

    except Exception as e:
        error_msg = str(e)
        if "Duplicate entry" in error_msg:
            messagebox.showerror("Ошибка", "Клиент с таким телефоном уже существует!")
        else:
            messagebox.showerror("Ошибка", f"Не удалось обновить клиента: {error_msg}")


def delete_client_from_db():
    """Удаление выбранного клиента"""
    selection = clients_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите клиента для удаления!")
        return

    item = clients_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 3:
        messagebox.showerror("Ошибка", "Не удалось получить данные клиента!")
        return

    client_id = values[0]  # Скрытый ID

    # Подтверждение удаления
    if not messagebox.askyesno("Подтверждение", "Вы уверены, что хотите удалить этого клиента?"):
        return

    # Удаляем клиента
    try:
        result = delete_client(client_id)
        if result > 0:
            messagebox.showinfo("Успех", "Клиент удален!")
            refresh_clients_page()

            # Очищаем форму
            client_name_entry.delete(0, tk.END)
            client_phone_entry.delete(0, tk.END)

            # Обновляем комбобокс на странице записи
            if hasattr(master_combo, 'winfo_exists'):
                refresh_booking_page()
        else:
            messagebox.showerror("Ошибка", "Не удалось удалить клиента!")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось удалить клиента: {str(e)}")


def add_client_from_form():
    """Добавление клиента из формы"""
    name = client_name_entry.get()
    phone = client_phone_entry.get()

    if not name or not phone:
        messagebox.showerror("Ошибка", "Заполните все поля!")
        return

    # Проверяем формат телефона
    if not validate_phone(phone):
        messagebox.showerror("Ошибка", "Неверный формат телефона! Должен быть: + и 11 цифр")
        return

    # Разбиваем имя на части
    name_parts = name.split()
    surname = name_parts[0] if len(name_parts) > 0 else ""
    first_name = name_parts[1] if len(name_parts) > 1 else name_parts[0] if name_parts else ""
    patronymic = name_parts[2] if len(name_parts) > 2 else ""

    # Добавляем клиента
    try:
        add_client(surname, first_name, patronymic, phone)
        messagebox.showinfo("Успех", "Клиент добавлен!")

        # Очищаем форму
        client_name_entry.delete(0, tk.END)
        client_phone_entry.delete(0, tk.END)

        # Обновляем таблицу
        refresh_clients_page()

        # Обновляем комбобокс на странице записи
        if hasattr(master_combo, 'winfo_exists'):
            refresh_booking_page()

    except Exception as e:
        error_msg = str(e)
        if "Duplicate entry" in error_msg:
            messagebox.showerror("Ошибка", "Клиент с таким телефоном уже существует!")
        else:
            messagebox.showerror("Ошибка", f"Не удалось добавить клиента: {error_msg}")


def create_clients_page(page):
    """Создание страницы клиентов"""
    page.grid_rowconfigure(2, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="Клиенты", font=LARGE_FONT).grid(row=0, column=0, pady=10, sticky="n")

    # Форма добавления/редактирования
    search_frame = ttk.Frame(page)
    search_frame.grid(row=1, column=0, sticky="ew", padx=12, pady=6)
    search_frame.grid_columnconfigure(0, weight=1)

    ttk.Label(search_frame, text="Имя:", font=MEDIUM_FONT).pack(side="left")
    global client_name_entry
    client_name_entry = ttk.Entry(search_frame, width=25)
    client_name_entry.pack(side="left", padx=6)
    make_large_entry(client_name_entry)

    ttk.Label(search_frame, text="Телефон:", font=MEDIUM_FONT).pack(side="left")
    global client_phone_entry
    client_phone_entry = ttk.Entry(search_frame, width=25)
    client_phone_entry.pack(side="left", padx=6)
    make_large_entry(client_phone_entry)

    # Кнопки формы
    form_buttons_frame = ttk.Frame(search_frame)
    form_buttons_frame.pack(side="left", padx=6)

    ttk.Button(form_buttons_frame, text="Добавить", style="Large.TButton",
               command=add_client_from_form).pack(side="left", padx=2)
    ttk.Button(form_buttons_frame, text="Очистить", style="Large.TButton",
               command=lambda: [client_name_entry.delete(0, tk.END),
                                client_phone_entry.delete(0, tk.END)]).pack(side="left", padx=2)

    # Таблица клиентов
    tree_frame = ttk.Frame(page)
    tree_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=(6, 0))
    tree_frame.grid_rowconfigure(0, weight=1)
    tree_frame.grid_columnconfigure(0, weight=1)

    # Добавляем столбец для ID (будет скрыт)
    columns = ("id", "name", "phone")
    global clients_tree
    clients_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=15)
    clients_tree.grid(row=0, column=0, sticky="nsew")

    # Настройка столбцов (первый скрыт)
    clients_tree.heading("name", text="Имя")
    clients_tree.heading("phone", text="Телефон")

    clients_tree.column("name", width=250)
    clients_tree.column("phone", width=150)

    # Привязываем обработчик выбора
    clients_tree.bind("<<TreeviewSelect>>", on_client_select)

    # Scrollbar
    scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=clients_tree.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    clients_tree.configure(yscrollcommand=scrollbar.set)

    # Кнопки действий под таблицей
    action_frame = ttk.Frame(page)
    action_frame.grid(row=3, column=0, pady=10, padx=12, sticky="ew")
    action_frame.grid_columnconfigure(0, weight=1)
    action_frame.grid_columnconfigure(1, weight=1)

    ttk.Button(action_frame, text="Изменить", style="Large.TButton",
               command=edit_client).grid(row=0, column=0, padx=4, sticky="ew")
    ttk.Button(action_frame, text="Удалить", style="Large.TButton",
               command=delete_client_from_db).grid(row=0, column=1, padx=4, sticky="ew")

    # Загружаем данные
    refresh_clients_page()


# ========== СТРАНИЦА МАСТЕРОВ ==========
def refresh_masters_page():
    """Обновление страницы мастеров"""
    # Очищаем таблицу
    for item in masters_tree.get_children():
        masters_tree.delete(item)

    # Загружаем мастеров из БД
    masters = get_masters()
    for master in masters:
        # Формируем полное имя
        full_name = f"{master['surname']} {master['name']}"
        if master.get('patronymic'):
            full_name += f" {master['patronymic']}"

        masters_tree.insert("", tk.END, values=(
            master['id'],  # Скрытый ID в первой колонке
            master['surname'],
            master['name'],
            master.get('patronymic', '') or '',
            full_name  # Полное имя для отображения
        ))

    # Скрываем первый (ID) и последний (полное имя) столбцы
    masters_tree.column("#1", width=0, stretch=False)
    masters_tree.heading("#1", text="")
    masters_tree.column("#5", width=0, stretch=False)  # Скрываем полное имя
    masters_tree.heading("#5", text="")


def on_master_select(event):
    """Обработка выбора мастера в таблице"""
    selection = masters_tree.selection()
    if selection:
        item = masters_tree.item(selection[0])
        values = item['values']

        # Заполняем поля формы данными выбранного мастера
        if len(values) >= 4:  # У нас 5 значений, но нам нужны первые 4 (без полного имени)
            master_surname_entry.delete(0, tk.END)
            master_surname_entry.insert(0, values[1])  # Фамилия

            master_name_entry.delete(0, tk.END)
            master_name_entry.insert(0, values[2])  # Имя

            master_patronymic_entry.delete(0, tk.END)
            master_patronymic_entry.insert(0, values[3])  # Отчество


def edit_master():
    """Редактирование выбранного мастера"""
    selection = masters_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите мастера для редактирования!")
        return

    item = masters_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 4:
        messagebox.showerror("Ошибка", "Не удалось получить данные мастера!")
        return

    master_id = values[0]

    # Получаем данные из формы
    surname = master_surname_entry.get()
    name = master_name_entry.get()
    patronymic = master_patronymic_entry.get()

    if not surname or not name:
        messagebox.showerror("Ошибка", "Заполните фамилию и имя!")
        return

    # Обновляем мастера
    try:
        result = update_master(master_id, surname, name, patronymic)
        if result > 0:
            messagebox.showinfo("Успех", "Мастер обновлен!")

            # Обновляем таблицу
            refresh_masters_page()

            # Очищаем форму
            master_surname_entry.delete(0, tk.END)
            master_name_entry.delete(0, tk.END)
            master_patronymic_entry.delete(0, tk.END)

            # Безопасное обновление других страниц
            update_other_pages_after_master_change()

        else:
            messagebox.showerror("Ошибка", "Не удалось обновить мастера!")

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось обновить мастера: {str(e)}")


def delete_master_from_db():
    """Удаление выбранного мастера"""
    selection = masters_tree.selection()
    if not selection:
        messagebox.showerror("Ошибка", "Выберите мастера для удаления!")
        return

    item = masters_tree.item(selection[0])
    values = item['values']

    if not values or len(values) < 4:
        messagebox.showerror("Ошибка", "Не удалось получить данные мастера!")
        return

    master_id = values[0]

    confirm_msg = (
        f"Вы уверены, что хотите удалить мастера {values[1]} {values[2]}?\n\n"
        f"ВНИМАНИЕ: При удалении мастера также удалятся:\n"
        f"• Все его услуги\n"
        f"• Все записи к нему\n"
        f"• Его расписание\n\n"
        f"Это действие нельзя отменить!"
    )

    if not messagebox.askyesno("Подтверждение удаления", confirm_msg):
        return

    try:
        result = delete_master(master_id)
        if result > 0:
            messagebox.showinfo("Успех", "Мастер удален!")

            refresh_masters_page()

            master_surname_entry.delete(0, tk.END)
            master_name_entry.delete(0, tk.END)
            master_patronymic_entry.delete(0, tk.END)

            update_other_pages_after_master_change()

        else:
            messagebox.showerror("Ошибка", "Не удалось удалить мастера!")

    except Exception as e:
        error_msg = str(e)
        if "foreign key constraint" in error_msg.lower():
            messagebox.showerror("Ошибка",
                                 "Невозможно удалить мастера, так как у него есть активные записи или услуги!\n"
                                 "Сначала удалите все записи и услуги этого мастера.")
        else:
            messagebox.showerror("Ошибка", f"Не удалось удалить мастера: {error_msg}")


def update_other_pages_after_master_change():
    """Безопасное обновление других страниц после изменения мастера"""
    try:
        if current_page == "Запись" and 'master_combo' in globals():
            refresh_booking_page()
    except:
        pass

    try:
        if current_page == "Услуги" and 'service_master_combo' in globals():
            refresh_services_page()
    except:
        pass

    try:
        if current_page == "График работы" and 'master_combo_sched' in globals():
            refresh_schedule_page()
    except:
        pass


def add_master_from_form():
    """Добавление мастера из формы"""
    surname = master_surname_entry.get()
    name = master_name_entry.get()
    patronymic = master_patronymic_entry.get()

    if not surname or not name:
        messagebox.showerror("Ошибка", "Заполните фамилию и имя!")
        return

    # Добавляем мастера
    try:
        add_master(surname, name, patronymic)
        messagebox.showinfo("Успех", "Мастер добавлен!")

        # Очищаем форму
        master_surname_entry.delete(0, tk.END)
        master_name_entry.delete(0, tk.END)
        master_patronymic_entry.delete(0, tk.END)

        # Обновляем таблицу на текущей странице
        refresh_masters_page()

        # Обновляем списки на других страницах (только если они существуют)
        try:
            if 'master_combo' in globals() and master_combo.winfo_exists():
                refresh_booking_page()
        except:
            pass

        try:
            if 'service_master_combo' in globals() and service_master_combo.winfo_exists():
                refresh_services_page()
        except:
            pass

        try:
            if 'master_combo_sched' in globals() and master_combo_sched.winfo_exists():
                refresh_schedule_page()
        except:
            pass

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось добавить мастера: {str(e)}")


def create_masters_page(page):
    """Создание страницы мастеров"""
    page.grid_rowconfigure(2, weight=1)
    page.grid_columnconfigure(0, weight=1)

    # Заголовок
    ttk.Label(page, text="Мастера", font=LARGE_FONT).grid(row=0, column=0, pady=10, sticky="n")

    # Форма добавления/редактирования
    add_frame = ttk.Frame(page, padding=8, relief="groove", borderwidth=1)
    add_frame.grid(row=1, column=0, sticky="nsew", padx=12, pady=6)
    add_frame.grid_columnconfigure(0, weight=0)
    add_frame.grid_columnconfigure(1, weight=1)

    ttk.Label(add_frame, text="Фамилия:", font=MEDIUM_FONT).grid(row=0, column=0, sticky="w", padx=4, pady=4)
    global master_surname_entry
    master_surname_entry = ttk.Entry(add_frame)
    master_surname_entry.grid(row=0, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(master_surname_entry)

    ttk.Label(add_frame, text="Имя:", font=MEDIUM_FONT).grid(row=1, column=0, sticky="w", padx=4, pady=4)
    global master_name_entry
    master_name_entry = ttk.Entry(add_frame)
    master_name_entry.grid(row=1, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(master_name_entry)

    ttk.Label(add_frame, text="Отчество:", font=MEDIUM_FONT).grid(row=2, column=0, sticky="w", padx=4, pady=4)
    global master_patronymic_entry
    master_patronymic_entry = ttk.Entry(add_frame)
    master_patronymic_entry.grid(row=2, column=1, sticky="ew", padx=4, pady=4)
    make_large_entry(master_patronymic_entry)

    # Кнопки формы
    form_buttons_frame = ttk.Frame(add_frame)
    form_buttons_frame.grid(row=3, column=0, columnspan=2, pady=6)

    ttk.Button(form_buttons_frame, text="Добавить", style="Large.TButton",
               command=add_master_from_form).pack(side="left", padx=4)
    ttk.Button(form_buttons_frame, text="Очистить", style="Large.TButton",
               command=lambda: [master_surname_entry.delete(0, tk.END),
                                master_name_entry.delete(0, tk.END),
                                master_patronymic_entry.delete(0, tk.END)]).pack(side="left", padx=4)

    # Таблица мастеров
    tree_frame = ttk.Frame(page)
    tree_frame.grid(row=2, column=0, sticky="nsew", padx=12, pady=(6, 0))
    tree_frame.grid_rowconfigure(0, weight=1)
    tree_frame.grid_columnconfigure(0, weight=1)

    # Столбцы таблицы: ID (скрыт), Фамилия, Имя, Отчество, Полное имя (скрыт)
    columns = ("id", "surname", "name", "patronymic", "full_name")
    global masters_tree
    masters_tree = ttk.Treeview(tree_frame, columns=columns, show="headings", height=15)
    masters_tree.grid(row=0, column=0, sticky="nsew")

    # Настройка видимых столбцов
    masters_tree.heading("surname", text="Фамилия")
    masters_tree.heading("name", text="Имя")
    masters_tree.heading("patronymic", text="Отчество")

    masters_tree.column("surname", width=150, anchor="w")
    masters_tree.column("name", width=150, anchor="w")
    masters_tree.column("patronymic", width=150, anchor="w")

    # Привязываем обработчик выбора
    masters_tree.bind("<<TreeviewSelect>>", on_master_select)

    # Scrollbar
    scrollbar = ttk.Scrollbar(tree_frame, orient="vertical", command=masters_tree.yview)
    scrollbar.grid(row=0, column=1, sticky="ns")
    masters_tree.configure(yscrollcommand=scrollbar.set)

    # Кнопки действий под таблицей
    action_frame = ttk.Frame(page)
    action_frame.grid(row=3, column=0, pady=10, padx=12, sticky="ew")
    action_frame.grid_columnconfigure(0, weight=1)
    action_frame.grid_columnconfigure(1, weight=1)

    ttk.Button(action_frame, text="Изменить", style="Large.TButton",
               command=edit_master).grid(row=0, column=0, padx=4, sticky="ew")
    ttk.Button(action_frame, text="Удалить", style="Large.TButton",
               command=delete_master_from_db).grid(row=0, column=1, padx=4, sticky="ew")

    # Загружаем данные
    refresh_masters_page()


# ========== СТРАНИЦА ОТЧЕТОВ ==========
def create_reports_page(page):
    """Создание страницы отчетов"""
    page.grid_rowconfigure(0, weight=1)
    page.grid_columnconfigure(0, weight=1)

    report_frame = ttk.Frame(page, padding=12)
    report_frame.grid(row=0, column=0, sticky="nsew")

    report_frame.grid_rowconfigure(4, weight=1)
    report_frame.grid_columnconfigure(0, weight=1)
    report_frame.grid_columnconfigure(1, weight=1)
    report_frame.grid_columnconfigure(2, weight=1)

    ttk.Label(report_frame, text="Отчеты", font=LARGE_FONT).grid(
        row=0, column=0, columnspan=3, pady=(0, 10), sticky="n"
    )

    ttk.Label(report_frame, text="Выберите период", font=MEDIUM_FONT).grid(
        row=1, column=0, sticky="w", padx=(0, 8)
    )

    period = tk.StringVar(value="day")
    rad_frame = ttk.Frame(report_frame)
    rad_frame.grid(row=1, column=1, columnspan=2, sticky="w")
    ttk.Radiobutton(rad_frame, text="день", value="day", variable=period).pack(side="left", padx=6)
    ttk.Radiobutton(rad_frame, text="месяц", value="month", variable=period).pack(side="left", padx=6)
    ttk.Radiobutton(rad_frame, text="год", value="year", variable=period).pack(side="left", padx=6)

    ttk.Separator(report_frame, orient="horizontal").grid(
        row=2, column=0, columnspan=3, sticky="ew", pady=10
    )

    ttk.Label(report_frame, text="Тип отчета", font=MEDIUM_FONT).grid(
        row=3, column=0, columnspan=3, sticky="w", pady=(0, 6)
    )

    report_types = [
        "Процент посещаемости",
        "Прибыль по мастерам",
        "Топ 5 самых популярных услуг за период",
        "Загруженность мастеров",
        "Количество записей по мастерам за период"
    ]

    types_frame = ttk.Frame(report_frame)
    types_frame.grid(row=4, column=0, columnspan=2, sticky="nsew", pady=(0, 8))
    types_frame.grid_columnconfigure(0, weight=1)

    for i, text in enumerate(report_types):
        ttk.Label(types_frame, text=text, font=SMALL_FONT).grid(
            row=i, column=0, sticky="w", pady=6, padx=(0, 10)
        )

    btns_frame = ttk.Frame(report_frame)
    btns_frame.grid(row=4, column=2, sticky="ns", padx=(8, 0))
    btns_frame.grid_columnconfigure(0, weight=1)

    for i in range(len(report_types)):
        ttk.Button(btns_frame, text="Сформировать", style="Large.TButton").grid(
            row=i, column=0, sticky="ew", pady=6, ipady=8
        )




#------------------------------------------------------------------------------------------------------
#------------------------------------------------------------------------------------------------------
#                       ОТЧЕТЫ


def generate_attendance_report():
    """Генерация отчета по посещаемости"""
    appointments = get_appointments()

    # Подсчитываем статистику
    total = len(appointments)
    if total == 0:
        return "Нет данных для отчета"

    completed = len([a for a in appointments if a.get('status') == 'Выполнено'])
    cancelled = len([a for a in appointments if a.get('status') == 'Отменено'])
    pending = len([a for a in appointments if a.get('status') == 'Ожидается'])
    in_progress = len([a for a in appointments if a.get('status') == 'В процессе'])

    attendance_rate = (completed / total * 100) if total > 0 else 0

    report = f"""
    ОТЧЕТ ПО ПОСЕЩАЕМОСТИ
    =====================

    Общее количество записей: {total}
    Выполнено: {completed} ({attendance_rate:.1f}%)
    Отменено: {cancelled}
    Ожидается: {pending}
    В процессе: {in_progress}

    Анализ:
    - Уровень посещаемости: {attendance_rate:.1f}%
    - Процент отмен: {(cancelled / total * 100):.1f}%
    """

    return report


def generate_master_profit_report():
    """Генерация отчета по прибыли по мастерам"""
    appointments = get_appointments()

    # Группируем по мастерам
    master_profit = {}
    for app in appointments:
        if app['status'] == 'Выполнено':
            master_name = app['master_name']
            profit = app.get('total_price', 0) or app.get('service_price', 0) or 0

            if master_name not in master_profit:
                master_profit[master_name] = {
                    'count': 0,
                    'profit': 0
                }

            master_profit[master_name]['count'] += 1
            master_profit[master_name]['profit'] += profit

    # Сортируем по прибыли
    sorted_masters = sorted(master_profit.items(), key=lambda x: x[1]['profit'], reverse=True)

    report = "ОТЧЕТ ПО ПРИБЫЛИ ПО МАСТЕРАМ\n"
    report += "============================\n\n"

    for master_name, data in sorted_masters:
        report += f"{master_name}:\n"
        report += f"  Количество выполненных услуг: {data['count']}\n"
        report += f"  Общая прибыль: {data['profit']} руб.\n"
        report += f"  Средняя прибыль за услугу: {(data['profit'] / data['count']):.0f} руб.\n\n"

    return report


def generate_top_services_report():
    """Генерация отчета по топ-5 услуг"""
    appointments = get_appointments()

    # Группируем по услугам
    service_count = {}
    for app in appointments:
        if app['status'] == 'Выполнено':
            service_name = app['service_name']
            service_count[service_name] = service_count.get(service_name, 0) + 1

    # Сортируем и берем топ-5
    top_services = sorted(service_count.items(), key=lambda x: x[1], reverse=True)[:5]

    report = "ТОП-5 САМЫХ ПОПУЛЯРНЫХ УСЛУГ\n"
    report += "============================\n\n"

    for i, (service_name, count) in enumerate(top_services, 1):
        report += f"{i}. {service_name}\n"
        report += f"   Количество выполненных: {count}\n\n"

    return report


def generate_master_load_report():
    """Генерация отчета по загруженности мастеров"""
    appointments = get_appointments()

    # Группируем по мастерам
    master_load = {}
    for app in appointments:
        if app['status'] != 'Отменено':
            master_name = app['master_name']
            if master_name not in master_load:
                master_load[master_name] = []
            master_load[master_name].append(app)

    report = "ОТЧЕТ ПО ЗАГРУЖЕННОСТИ МАСТЕРОВ\n"
    report += "===============================\n\n"

    for master_name, apps in master_load.items():
        report += f"{master_name}:\n"
        report += f"  Всего записей: {len(apps)}\n"

        # Группируем по статусам
        status_count = {}
        for app in apps:
            status = app['status']
            status_count[status] = status_count.get(status, 0) + 1

        for status, count in status_count.items():
            report += f"  {status}: {count}\n"

        report += "\n"

    return report


def generate_appointments_by_master_report():
    """Генерация отчета по количеству записей по мастерам"""
    appointments = get_appointments()

    # Группируем по мастерам и датам
    master_appointments = {}
    for app in appointments:
        master_name = app['master_name']
        app_date = app['appointment_date']

        if master_name not in master_appointments:
            master_appointments[master_name] = {}

        if app_date not in master_appointments[master_name]:
            master_appointments[master_name][app_date] = 0

        master_appointments[master_name][app_date] += 1

    report = "КОЛИЧЕСТВО ЗАПИСЕЙ ПО МАСТЕРАМ ЗА ПЕРИОД\n"
    report += "======================================\n\n"

    for master_name, dates in master_appointments.items():
        report += f"{master_name}:\n"
        total = 0
        for app_date, count in dates.items():
            report += f"  {app_date.strftime('%d.%m.%Y')}: {count} записей\n"
            total += count
        report += f"  Итого: {total} записей\n\n"

    return report


def save_report_to_docx(report_text, report_type):
    """Сохранение отчета в файл .docx"""
    # Открываем диалог выбора файла
    file_path = filedialog.asksaveasfilename(
        defaultextension=".docx",
        filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")],
        initialfile=f"{report_type}.docx"
    )

    if not file_path:
        return False  # Пользователь отменил сохранение

    try:
        # Создаем документ
        doc = Document()

        # Добавляем заголовок
        title = doc.add_heading(report_type, 0)
        title.alignment = 1  # Центрирование

        # Добавляем дату
        from datetime import datetime
        date_para = doc.add_paragraph()
        date_para.add_run(f"Сформировано: {datetime.now().strftime('%d.%m.%Y %H:%M')}").italic = True
        date_para.alignment = 1

        doc.add_paragraph()  # Пустая строка

        # Разбиваем текст на параграфы и добавляем в документ
        lines = report_text.split('\n')
        for line in lines:
            if line.strip() == '':
                doc.add_paragraph()  # Пустая строка
            elif line.startswith('='):
                # Заголовок второго уровня
                heading = doc.add_heading(line.strip('= '), level=2)
            elif ':' in line and not line.startswith('  '):
                # Важные пункты
                p = doc.add_paragraph()
                p.add_run(line).bold = True
            else:
                # Обычный текст
                doc.add_paragraph(line)

        # Сохраняем документ
        doc.save(file_path)

        # Открываем файл (опционально)
        if os.name == 'nt':  # Windows
            os.startfile(file_path)
        elif os.name == 'posix':  # Linux, Mac
            os.system(f'open "{file_path}"' if sys.platform == 'darwin' else f'xdg-open "{file_path}"')

        return True

    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось сохранить отчет: {str(e)}")
        return False


def show_report_dialog(report_type, report_function):
    """Показать диалоговое окно с отчетом"""
    # Создаем новое окно
    report_window = tk.Toplevel(root)
    report_window.title(f"Отчет: {report_type}")
    report_window.geometry("800x600")
    report_window.transient(root)  # Сделать окно дочерним

    # Создаем текстовое поле для отчета
    text_frame = ttk.Frame(report_window)
    text_frame.pack(fill="both", expand=True, padx=10, pady=10)

    text_widget = tk.Text(text_frame, wrap="word", font=("Courier", 10))
    scrollbar = ttk.Scrollbar(text_frame, orient="vertical", command=text_widget.yview)
    text_widget.configure(yscrollcommand=scrollbar.set)

    text_widget.pack(side="left", fill="both", expand=True)
    scrollbar.pack(side="right", fill="y")

    # Генерируем отчет
    report_text = report_function()
    text_widget.insert("1.0", report_text)
    text_widget.configure(state="disabled")  # Запрещаем редактирование

    # Фрейм с кнопками
    button_frame = ttk.Frame(report_window)
    button_frame.pack(fill="x", padx=10, pady=(0, 10))

    # Кнопка сохранения
    def save_report():
        if save_report_to_docx(report_text, report_type):
            messagebox.showinfo("Успех", f"Отчет сохранен!")

    ttk.Button(button_frame, text="Сохранить в Word",
               command=save_report, style="Large.TButton").pack(side="right", padx=5)

    # Кнопка закрытия
    ttk.Button(button_frame, text="Закрыть",
               command=report_window.destroy).pack(side="right", padx=5)


def create_reports_page(page):
    """Создание страницы отчетов"""
    page.grid_rowconfigure(0, weight=1)
    page.grid_columnconfigure(0, weight=1)

    report_frame = ttk.Frame(page, padding=12)
    report_frame.grid(row=0, column=0, sticky="nsew")

    report_frame.grid_rowconfigure(4, weight=1)
    report_frame.grid_columnconfigure(0, weight=1)
    report_frame.grid_columnconfigure(1, weight=1)
    report_frame.grid_columnconfigure(2, weight=1)

    ttk.Label(report_frame, text="Отчеты", font=LARGE_FONT).grid(
        row=0, column=0, columnspan=3, pady=(0, 10), sticky="n"
    )

    ttk.Label(report_frame, text="Выберите период", font=MEDIUM_FONT).grid(
        row=1, column=0, sticky="w", padx=(0, 8)
    )

    period = tk.StringVar(value="day")
    rad_frame = ttk.Frame(report_frame)
    rad_frame.grid(row=1, column=1, columnspan=2, sticky="w")
    ttk.Radiobutton(rad_frame, text="день", value="day", variable=period).pack(side="left", padx=6)
    ttk.Radiobutton(rad_frame, text="месяц", value="month", variable=period).pack(side="left", padx=6)
    ttk.Radiobutton(rad_frame, text="год", value="year", variable=period).pack(side="left", padx=6)

    ttk.Separator(report_frame, orient="horizontal").grid(
        row=2, column=0, columnspan=3, sticky="ew", pady=10
    )

    ttk.Label(report_frame, text="Тип отчета", font=MEDIUM_FONT).grid(
        row=3, column=0, columnspan=3, sticky="w", pady=(0, 6)
    )

    report_types = [
        ("Процент посещаемости", generate_attendance_report),
        ("Прибыль по мастерам", generate_master_profit_report),
        ("Топ 5 самых популярных услуг", generate_top_services_report),
        ("Загруженность мастеров", generate_master_load_report),
        ("Количество записей по мастерам", generate_appointments_by_master_report)
    ]

    types_frame = ttk.Frame(report_frame)
    types_frame.grid(row=4, column=0, columnspan=2, sticky="nsew", pady=(0, 8))
    types_frame.grid_columnconfigure(0, weight=1)

    btns_frame = ttk.Frame(report_frame)
    btns_frame.grid(row=4, column=2, sticky="ns", padx=(8, 0))
    btns_frame.grid_columnconfigure(0, weight=1)

    for i, (text, func) in enumerate(report_types):
        # Метка с названием отчета
        ttk.Label(types_frame, text=text, font=SMALL_FONT).grid(
            row=i, column=0, sticky="w", pady=6, padx=(0, 10)
        )

        # Кнопка для формирования отчета
        ttk.Button(btns_frame, text="Сформировать", style="Large.TButton",
                   command=lambda f=func, t=text: show_report_dialog(t, f)).grid(
            row=i, column=0, sticky="ew", pady=6, ipady=8
        )





# ========== ОСНОВНОЙ ИНТЕРФЕЙС ==========
def main():
    """Основная функция приложения"""
    global root

    # Создаем главный контейнер
    main_container = ttk.Frame(root)
    main_container.pack(fill="both", expand=True)

    # Левое меню
    menu_frame = ttk.Frame(main_container, width=260)
    menu_frame.pack(side="left", fill="y")
    menu_frame.pack_propagate(False)

    # Область контента
    content_frame = ttk.Frame(main_container)
    content_frame.pack(side="right", fill="both", expand=True)

    # Создаем кнопки меню
    for name in page_names:
        btn = ttk.Button(menu_frame, text=name, style="Menu.TButton",
                         command=lambda n=name: show_page(n))
        btn.pack(fill="x", pady=4, ipady=8, expand=True)
        menu_buttons[name] = btn

    # Кнопка выхода
    ttk.Button(menu_frame, text="Выход", style="Menu.TButton",
               command=quit_app).pack(fill="x", pady=(8, 8), ipady=8)

    # Создаем страницы
    for name in page_names:
        frame = ttk.Frame(content_frame)
        frame.grid(row=0, column=0, sticky="nsew")
        pages[name] = frame

    content_frame.grid_rowconfigure(0, weight=1)
    content_frame.grid_columnconfigure(0, weight=1)

    # Настраиваем стиль таблиц
    style.configure("Treeview", rowheight=40)

    # Создаем все страницы
    create_main_page(pages["Главная"])
    create_booking_page(pages["Запись"])
    create_schedule_page(pages["График работы"])
    create_services_page(pages["Услуги"])
    create_clients_page(pages["Клиенты"])
    create_masters_page(pages["Мастера"])
    create_reports_page(pages["Отчеты"])

    # Показываем главную страницу
    show_page("Главная")

    # Запускаем главный цикл
    root.mainloop()


if __name__ == "__main__":
    main()